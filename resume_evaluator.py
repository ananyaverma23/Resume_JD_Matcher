import json
import os
import time
from pathlib import Path

from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document


# ============================================================
# 1. ENVIRONMENT + GROQ CLIENT
# ============================================================

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("GROQ_API_KEY environment variable is not set.")

client = Groq(api_key=my_api_key)

# Use the larger model for better extraction and matching quality.
model = "openai/gpt-oss-120b"


# ============================================================
# 2. JOB DESCRIPTION
# ============================================================

job_description = """
Description
Do you want to solve real customer problems through innovative technology? Do you enjoy working on scalable services in a collaborative team environment? Do you want to see your code directly impact millions of customers worldwide?

At Amazon, we hire the best minds in technology to innovate and build on behalf of our customers. Customer obsession is part of our company DNA, which has made us one of the world's most beloved brands.

Our Software Development Engineers (SDEs) use modern technology to solve complex problems while seeing their work's impact first-hand. The challenges SDEs solve at Amazon are meaningful and influence millions of customers, sellers, and products globally. We seek individuals passionate about creating new products, features, and services while managing ambiguity in an environment where development cycles are measured in weeks, not years.

At Amazon, we believe in ownership at every level. As an SDE-I, you'll own the entire lifecycle of your code - from design through deployment and ongoing operations. This ownership mindset, combined with our commitment to operational excellence, ensures we deliver the highest quality solutions for our customers.

We're looking for curious minds who think big and want to define tomorrow's technology. At Amazon, you'll grow into the high-impact engineer you know you can be, supported by a culture of learning and mentorship. Every day brings exciting new challenges and opportunities for personal growth.

Key job responsibilities
• Collaborate and communicate effectively with experienced cross-disciplinary Amazonians to design, build, and operate innovative products and services that delight our customers, while participating in technical discussions to drive solutions forward.
• Design and develop scalable solutions using cloud-native architectures and microservices in a large distributed computing environment.
• Participate in code reviews and contribute to technical documentation.
• Build and maintain resilient distributed systems that are scalable, fault-tolerant, and cost-effective.
• Leverage and contribute to the development of GenAI and AI-powered tools to enhance development productivity while staying current with emerging technologies.
• Write clean, maintainable code following best practices and design patterns.
• Work in an agile environment practicing CI/CD principles while participating in operational responsibilities including on-call duties.
• Demonstrate operational excellence through monitoring, troubleshooting, and resolving production issues.

Basic Qualifications
- Experience with at least one general-purpose programming language such as Java, Python, C++, C#, Go, Rust, or TypeScript
- Experience with data structure implementation, basic algorithm development, and/or object-oriented design principles
- Currently has, or is in the process of obtaining a bachelor’s degree in Computer Science, Computer Engineering, Data Science, Information Systems, or related STEM fields
- Must be 18 years of age of older

Preferred Qualifications
- Experience from previous technical internship(s) or demonstrated project experience
- Experience with one or more of the following: AI tools for development productivity, Cloud platforms (preferably AWS), Database systems (SQL and NoSQL), Contributing to open-source projects, Version control systems, Debugging and troubleshooting complex systems
- Demonstrated ability to learn and adapt to new technologies quickly
- Basic understanding of software development lifecycle (SDLC)
- Strong problem-solving and analytical skills
- Excellent written and verbal communication skills
"""


# ============================================================
# 3. PYDANTIC SCHEMAS
# ============================================================

class JobDescription(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = Field(default_factory=list)


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    total_experience_years: float | None = None

    skills: list[str] = Field(default_factory=list)
    experiences: list[Experience] = Field(default_factory=list)
    education: list[str] = Field(default_factory=list)
    projects: list[str] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)


class MatchResult(BaseModel):
    score: float
    details: dict


# Generate JSON schemas that are shown to the LLM.
job_schema = JobDescription.model_json_schema()
resume_schema = Resume.model_json_schema()
match_schema = MatchResult.model_json_schema()


# ============================================================
# 4. PARSE JOB DESCRIPTION
# ============================================================

def parse_job_description(job_text: str) -> JobDescription:
    system_prompt = f"""
You are an expert HR assistant.

Your job is to analyze job descriptions and extract
structured information from them.

Return ONLY valid JSON matching this schema:

{job_schema}

IMPORTANT:
- Do NOT return the schema itself.
- Do NOT return fields like "properties", "title", or "type".
- Fill the schema with actual information extracted from the job description.
- If minimum experience is not mentioned, return null.
- If information for a list is missing, return an empty list.
- Do not invent information.
"""

    user_prompt = f"""
Analyze the following job description:

{job_text}
"""

    messages = [
        {
            "role": "system",
            "content": system_prompt
        },
        {
            "role": "user",
            "content": user_prompt
        }
    ]

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format={"type": "json_object"}
    )

    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)

    return JobDescription(**data)


# ============================================================
# 5. PARSE RESUME
# ============================================================

def parse_resume(resume_text: str) -> Resume:
    system_prompt = f"""
You are an expert resume parser.

Extract information from the resume based on its meaning,
not only exact section headings.

For example, these can all represent experience:
- Experience
- Professional Experience
- Work History
- Employment
- Internships

Skills may appear in the skills section, work experience,
internships, or projects.

Return ONLY valid JSON matching this schema:

{resume_schema}

IMPORTANT RULES:
1. Extract only information explicitly present in the resume.
2. Do not invent information.
3. If a value is not available, return null.
4. If a list has no information, return an empty list.
5. Include internships inside experiences.
6. Extract relevant skills from across the entire resume.
"""

    user_prompt = f"""
Parse the following resume:

{resume_text}
"""

    messages = [
        {
            "role": "system",
            "content": system_prompt
        },
        {
            "role": "user",
            "content": user_prompt
        }
    ]

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format={"type": "json_object"}
    )

    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)

    return Resume(**data)


# ============================================================
# 6. MATCH RESUME WITH JOB DESCRIPTION
# ============================================================

def calculate_match(job: JobDescription, resume: Resume) -> MatchResult:
    prompt = f"""
You are an HR recruiter.

Compare the candidate's resume with the job description.

JOB DESCRIPTION:
{job.model_dump_json(indent=2)}

CANDIDATE RESUME:
{resume.model_dump_json(indent=2)}

Return JSON matching this schema:

{match_schema}

Give me:
1. Candidate name
2. Matching skills
3. Missing important skills
4. Whether experience requirement is met
5. Overall match percentage from 0 to 100
6. A short final verdict

Keep the response concise and easy to read.
"""

    messages = [
        {
            "role": "user",
            "content": prompt
        }
    ]

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format={"type": "json_object"}
    )

    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)

    return MatchResult(**data)


# ============================================================
# 7. READ PDF
# ============================================================

def read_pdf(file_path: Path) -> str:
    reader = PdfReader(file_path)
    text = ""

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


# ============================================================
# 8. READ DOCX
# ============================================================

def read_docx(file_path: Path) -> str:
    document = Document(file_path)
    text = ""

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"

    return text


# ============================================================
# 9. READ RESUME BASED ON FILE TYPE
# ============================================================

def read_resume(file_path: Path) -> str | None:
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return read_pdf(file_path)

    if extension == ".docx":
        return read_docx(file_path)

    return None


# ============================================================
# 10. MAIN PIPELINE
# ============================================================

job = parse_job_description(job_description)

print("Minimum experience:", job.minimum_experience)
print("Education requirements:", job.education_requirements)

resume_folder = Path("resumes")
all_results = []

for file_path in resume_folder.iterdir():

    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue

    print("\nProcessing:", file_path.name)

    resume_text = read_resume(file_path)

    if not resume_text:
        print("Could not extract resume text.")
        continue

    # Useful while debugging PDF/DOCX extraction.
    print("Resume characters:", len(resume_text))
    print("Resume preview:")
    print(resume_text[:2000])

    parsed_resume = parse_resume(resume_text)

    # Avoid sending requests too quickly.
    time.sleep(5)

    result = calculate_match(job, parsed_resume)

    print("Score:", result.score)

    all_results.append(
        {
            "name": parsed_resume.name,
            "score": result.score,
            "details": result.details
        }
    )

    # Avoid sending requests too quickly.
    time.sleep(5)


# ============================================================
# 11. RANK ALL CANDIDATES
# ============================================================

all_results.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)

top_2 = all_results[:2]
worst_2 = all_results[-2:]


# ============================================================
# 12. DISPLAY RESULTS
# ============================================================

print("\nTOP 2 CANDIDATES")
for idx, candidate in enumerate(top_2, start=1):
    details = candidate["details"]
    print(f"{idx}. {candidate['name']} - {candidate['score']}%")
    print("   Matching skills:", details.get("matching_skills", []))
    print("   Missing important skills:", details.get("missing_important_skills", details.get("missing_skills", [])))
    print("   Experience requirement met:", details.get("experience_requirement_met", details.get("experience_met", "N/A")))
    print("   Overall match percentage:", details.get("overall_match_percentage", details.get("overall_match", "N/A")))
    print("   Verdict:", details.get("final_verdict", details.get("verdict", "N/A")))

print("\nLOWEST 2 CANDIDATES")
for idx, candidate in enumerate(worst_2, start=1):
    details = candidate["details"]
    print(f"{idx}. {candidate['name']} - {candidate['score']}%")
    print("   Matching skills:", details.get("matching_skills", []))
    print("   Missing important skills:", details.get("missing_important_skills", details.get("missing_skills", [])))
    print("   Experience requirement met:", details.get("experience_requirement_met", details.get("experience_met", "N/A")))
    print("   Overall match percentage:", details.get("overall_match_percentage", details.get("overall_match", "N/A")))
    print("   Verdict:", details.get("final_verdict", details.get("verdict", "N/A")))